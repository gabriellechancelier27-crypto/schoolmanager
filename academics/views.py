from django.shortcuts import render

# Create your views here.
import openpyxl
from django.http import HttpResponse
from PIL import Image
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from core.mixins import SchoolScopedMixin
from core.permissions import ReadOnlyOrManager
from core.services import log_activity

from .models import (ClassRoom, Enrollment, Parent, Student, Subject,
                     Teacher, TeachingAssignment)
from .serializers import (ClassRoomSerializer, EnrollmentSerializer, ParentSerializer,
                          StudentSerializer, StudentWriteSerializer, SubjectSerializer,
                          TeacherSerializer, TeachingAssignmentSerializer)


class TeacherViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = Teacher.objects.all()
    serializer_class = TeacherSerializer
    permission_classes = [ReadOnlyOrManager]
    filterset_fields = ["status", "gender"]
    search_fields = ["matricule", "first_name", "last_name", "phone"]


class ParentViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = Parent.objects.all()
    serializer_class = ParentSerializer
    permission_classes = [ReadOnlyOrManager]
    search_fields = ["first_name", "last_name", "phone", "email"]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if user.role == "PARENT":
            return qs.filter(user=user)   # un parent ne voit que sa propre fiche
        if user.role == "STUDENT":
            return qs.none()
        return qs


class SubjectViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = Subject.objects.all()
    serializer_class = SubjectSerializer
    permission_classes = [ReadOnlyOrManager]
    search_fields = ["name", "code"]


class ClassRoomViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = ClassRoom.objects.select_related("main_teacher")
    serializer_class = ClassRoomSerializer
    permission_classes = [ReadOnlyOrManager]
    filterset_fields = ["level", "is_active"]
    search_fields = ["name", "level", "serie"]


class StudentViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = Student.objects.prefetch_related("student_parents__parent")
    permission_classes = [ReadOnlyOrManager]
    filterset_fields = ["status", "gender"]
    search_fields = ["matricule", "first_name", "last_name", "phone", "birth_place"]

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return StudentWriteSerializer
        return StudentSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        class_id = self.request.query_params.get("classe")
        if class_id:
            qs = qs.filter(enrollments__class_room_id=class_id,
                           enrollments__status=Enrollment.Status.ACTIVE).distinct()
        if user.role == "PARENT":
            qs = qs.filter(student_parents__parent__user=user).distinct()
        elif user.role == "STUDENT":
            qs = qs.filter(user=user)
        return qs

    @action(detail=True, methods=["post"], url_path="photo",
            parser_classes=[MultiPartParser, FormParser],
            permission_classes=[ReadOnlyOrManager])
    def upload_photo(self, request, pk=None):
        """Envoi de la photo d'un élève (image, 2 Mo maximum)."""
        student = self.get_object()
        photo = request.FILES.get("photo")
        if photo is None:
            return Response({"detail": "Fichier manquant (champ « photo »)."}, status=400)
        if photo.size > 2 * 1024 * 1024:
            return Response({"detail": "Image trop lourde (2 Mo maximum)."}, status=400)
        try:
            Image.open(photo).verify()
        except Exception:
            return Response({"detail": "Le fichier n'est pas une image valide."}, status=400)
        photo.seek(0)
        student.photo = photo
        student.save()
        return Response(StudentSerializer(student, context={"request": request}).data)

    @action(detail=False, methods=["get"], permission_classes=[ReadOnlyOrManager])
    def export(self, request):
        """Export Excel de la liste d'élèves (§4.2)."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Élèves"
        sheet.append(["Matricule", "Nom", "Prénom", "Sexe", "Date naissance", "Lieu",
                      "Téléphone", "Statut", "Classe actuelle"])
        for s in self.filter_queryset(self.get_queryset()):
            sheet.append([s.matricule, s.last_name, s.first_name, s.gender,
                          str(s.birth_date or ""), s.birth_place, s.phone,
                          s.get_status_display(), s.current_class_name or ""])
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="eleves.xlsx"'
        workbook.save(response)
        return response

    @action(detail=False, methods=["post"], url_path="import",
            permission_classes=[ReadOnlyOrManager])
    def import_excel(self, request):
        """Import Excel : Matricule | Nom | Prénom | Sexe (matricule facultatif)."""
        user = request.user
        if user.school is None:
            return Response({"detail": "Compte non rattaché à un établissement."}, status=400)
        uploaded = request.FILES.get("file")
        if uploaded is None:
            return Response({"detail": "Fichier manquant (champ « file »)."}, status=400)
        workbook = openpyxl.load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        created, updated, errors = 0, 0, []
        for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not row or len(row) < 2 or not row[1]:
                continue
            matricule = str(row[0] or "").strip()
            last_name = str(row[1]).strip()
            first_name = str(row[2] or "").strip() if len(row) > 2 else ""
            gender = (str(row[3] or "M").strip()[:1].upper() if len(row) > 3 else "M")
            if gender not in ("M", "F"):
                gender = "M"
            try:
                if matricule:
                    _, was_created = Student.objects.update_or_create(
                        school=user.school, matricule=matricule,
                        defaults={"last_name": last_name, "first_name": first_name,
                                  "gender": gender})
                    created += int(was_created)
                    updated += int(not was_created)
                else:
                    Student.objects.create(school=user.school, last_name=last_name,
                                           first_name=first_name, gender=gender)
                    created += 1
            except Exception as exc:
                errors.append(f"Ligne {index} : {exc}")
        return Response({"crees": created, "mis_a_jour": updated, "erreurs": errors})


class EnrollmentViewSet(viewsets.ModelViewSet):
    queryset = Enrollment.objects.select_related("student", "class_room", "school_year")
    serializer_class = EnrollmentSerializer
    permission_classes = [ReadOnlyOrManager]
    filterset_fields = ["status", "class_room", "school_year"]
    search_fields = ["student__matricule", "student__first_name", "student__last_name"]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if user.is_super_admin:
            return qs
        if not user.school_id:
            return qs.none()
        qs = qs.filter(student__school=user.school)
        if user.role == "TEACHER" and hasattr(user, "teacher_profile"):
            qs = qs.filter(class_room__assignments__teacher=user.teacher_profile).distinct()
        elif user.role == "PARENT":
            qs = qs.filter(student__student_parents__parent__user=user).distinct()
        elif user.role == "STUDENT":
            qs = qs.filter(student__user=user)
        return qs

    def _check_school(self, serializer):
        user = self.request.user
        if not user.is_super_admin:
            for key in ("student", "class_room", "school_year"):
                obj = serializer.validated_data.get(key)
                if obj is not None and getattr(obj, "school_id", None) not in (None, user.school_id):
                    raise PermissionDenied("Inscription impossible hors de votre établissement.")

    def perform_create(self, serializer):
        self._check_school(serializer)
        serializer.save()

    def perform_update(self, serializer):
        self._check_school(serializer)
        serializer.save()

    def perform_destroy(self, instance):
        # On ne supprime jamais une inscription qui porte des notes ou des paiements
        if instance.grades.exists() or instance.payments.exists():
            raise ValidationError({"detail": (
                "Suppression impossible : des notes ou des paiements sont liés à cette "
                "inscription. Passez-la plutôt au statut « Annulée ».")})
        log_activity(self.request.user, "DELETE", instance, f"Suppression de l'inscription {instance}")
        instance.delete()


class TeachingAssignmentViewSet(SchoolScopedMixin, viewsets.ModelViewSet):
    queryset = TeachingAssignment.objects.select_related("teacher", "subject",
                                                         "class_room", "school_year")
    serializer_class = TeachingAssignmentSerializer
    permission_classes = [ReadOnlyOrManager]
    filterset_fields = ["school_year", "class_room", "subject", "teacher"]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if user.role == "TEACHER" and hasattr(user, "teacher_profile"):
            qs = qs.filter(teacher=user.teacher_profile)  # l'enseignant voit ses affectations
        return qs